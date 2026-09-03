import openpyxl, json, sys

def dump(path, name):
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"\n{'='*60}\n📄 {name}  (sheets: {wb.sheetnames})\n{'='*60}")
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n--- Sheet: {sn} (rows={ws.max_row} cols={ws.max_column}) ---")
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True):
            if any(c is not None and str(c).strip() for c in row):
                print(" | ".join("" if c is None else str(c).strip() for c in row))

dump(r"C:\Users\Administrator\Desktop\课程表.xlsx", "课程表")
dump(r"C:\Users\Administrator\Desktop\作息表.xlsx", "作息表")
